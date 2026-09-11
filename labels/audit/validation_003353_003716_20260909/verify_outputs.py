import json
import hashlib
from copy import copy
from pathlib import Path
import openpyxl

OUT=Path(__file__).resolve().parent
LABELS=OUT.parents[1]
manifest=json.loads((OUT/'manifest.json').read_text(encoding='utf-8'))
data=json.loads((OUT/'workbook_data.json').read_text(encoding='utf-8'))
original=openpyxl.load_workbook(LABELS/'validation_003353_003716_manual.xlsx',data_only=False)
source=original['검수 정답지']
results={}
def blank(v): return '' if v is None else v
for stage in ('baseline','p3'):
    path=LABELS/f'validation_003353_003716_{stage}_20260909.xlsx'
    wb=openpyxl.load_workbook(path,data_only=False)
    values=openpyxl.load_workbook(path,data_only=True)
    sheet,cached=wb['검수 정답지'],values['검수 정답지']
    assert wb.sheetnames==original.sheetnames
    assert sheet.max_row==387 and sheet.max_column==8
    assert [sheet.cell(1,c).value for c in range(1,9)]==manifest['headers']
    for row,expected in enumerate(data[stage]['rows'],2):
        for col in range(1,9):
            assert blank(cached.cell(row,col).value)==blank(expected[col-1]),(stage,row,col,cached.cell(row,col).value,expected[col-1])
        for col in (1,3,4):
            assert sheet.cell(row,col).value==source.cell(row,col).value
    assert list(sheet.tables)==list(source.tables)
    for name in source.tables:
        assert sheet.tables[name].ref==source.tables[name].ref
    assert len(sheet.data_validations.dataValidation)==len(source.data_validations.dataValidation)
    assert len(sheet.conditional_formatting)==len(source.conditional_formatting)
    assert sheet.sheet_view.pane.ySplit==source.sheet_view.pane.ySplit
    assert sheet.sheet_view.pane.state==source.sheet_view.pane.state
    styles=[]
    for row in range(1,388):
        for col in range(1,9):
            a,b=source.cell(row,col),sheet.cell(row,col)
            for attr in ('font','fill','border','alignment','number_format','protection'):
                if copy(getattr(a,attr))!=copy(getattr(b,attr)):
                    styles.append({'cell':a.coordinate,'attribute':attr})
    results[stage]={'rows':386,'values_and_formulas_match':True,'tables_and_panes_preserved':True,
                    'style_differences':styles[:30],'style_difference_count':len(styles),
                    'sha256':hashlib.sha256(path.read_bytes()).hexdigest()}
for path,sha in manifest['sources'].items():
    assert hashlib.sha256(Path(path).read_bytes()).hexdigest()==sha
for stage in ('baseline','p3'):
    meta=json.loads((OUT/f'{stage}_runtime.json').read_text(encoding='utf-8'))
    for path,sha in meta['code_sha256'].items():
        assert hashlib.sha256(Path(path).read_bytes()).hexdigest()==sha,'Production source changed'
results['original_sources_unchanged']=True
(OUT/'output_verification.json').write_text(json.dumps(results,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(results,ensure_ascii=False,indent=2))
