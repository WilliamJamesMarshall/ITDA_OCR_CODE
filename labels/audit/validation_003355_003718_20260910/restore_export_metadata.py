"""Restore font classification and empty diagonal metadata omitted by export.

All data/formulas are authored and calculated by Artifact Tool. This restores
only unsupported OOXML style metadata in newly created result files.
Openpyxl is used for read-only style-ID mapping, never for workbook authoring.
"""
import hashlib
import json
from copy import deepcopy
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from xml.etree import ElementTree as ET
import openpyxl

OUT=Path(__file__).resolve().parent
LABELS=OUT.parents[1]
source=LABELS/'validation_003355_003718_manual.xlsx'
ns={'m':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
ET.register_namespace('',ns['m'])
original=openpyxl.load_workbook(source)
with ZipFile(source) as z:
    old=ET.fromstring(z.read('xl/styles.xml'))
oldfonts=old.find('m:fonts',ns)
oldborders=old.find('m:borders',ns)
for stage in ('baseline','p3'):
    target=LABELS/f'validation_003355_003718_{stage}_20260910.xlsx'
    result=openpyxl.load_workbook(target)
    with ZipFile(target) as z:
        entries={info.filename:(info,z.read(info.filename)) for info in z.infolist()}
    styles=ET.fromstring(entries['xl/styles.xml'][1])
    fonts=styles.find('m:fonts',ns)
    borders=styles.find('m:borders',ns)
    mappings={'font':{},'border':{}}
    for row in range(1,388):
        for col in range(1,9):
            a=original.active.cell(row,col)._style
            b=result.active.cell(row,col)._style
            for kind,key in [('font','fontId'),('border','borderId')]:
                mappings[kind].setdefault(getattr(b,key),set()).add(getattr(a,key))
    for new_id,old_ids in mappings['font'].items():
        for tag in ('charset','family'):
            nodes=[oldfonts[i].find('m:'+tag,ns) for i in old_ids]
            assert len({ET.tostring(n) if n is not None else None for n in nodes})==1
            if fonts[new_id].find('m:'+tag,ns) is None and nodes[0] is not None:
                fonts[new_id].append(deepcopy(nodes[0]))
    for new_id,old_ids in mappings['border'].items():
        nodes=[oldborders[i].find('m:diagonal',ns) for i in old_ids]
        assert len({ET.tostring(n) if n is not None else None for n in nodes})==1
        if borders[new_id].find('m:diagonal',ns) is None and nodes[0] is not None:
            assert len(nodes[0])==0 and not nodes[0].attrib
            borders[new_id].append(deepcopy(nodes[0]))
    temporary=OUT/f'{stage}_metadata_restore.tmp.xlsx'
    assert not temporary.exists()
    with ZipFile(temporary,'x',compression=ZIP_DEFLATED) as z:
        for name,(info,content) in entries.items():
            z.writestr(info,ET.tostring(styles,encoding='utf-8',xml_declaration=True) if name=='xl/styles.xml' else content)
    temporary.replace(target)
    inspect=Path(str(target)+'.inspect.ndjson')
    lines=inspect.read_text(encoding='utf-8').splitlines()
    header=json.loads(lines[0])
    header['sha256']=hashlib.sha256(target.read_bytes()).hexdigest()
    lines[0]=json.dumps(header,ensure_ascii=False,separators=(',',':'))
    inspect.write_text('\n'.join(lines)+'\n',encoding='utf-8')
    print(stage, 'restored omitted font and border metadata')
