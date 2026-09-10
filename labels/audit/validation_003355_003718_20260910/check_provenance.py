"""Read-only source checks, recorded separately from image-only inference."""
import hashlib
import json
from pathlib import Path
from datetime import datetime, timezone
import openpyxl

OUT = Path(__file__).resolve().parent
OLD = OUT.with_name('validation_003355_003718_20260909')
ROOT = OUT.parents[2]
def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()
manifest = json.loads((OUT/'manifest.json').read_text(encoding='utf-8'))
checks = {p:digest(p)==sha for p,sha in manifest['sources'].items()}
checks.update({v['path']:digest(v['path'])==v['sha256'] for v in manifest['images'].values()})
runtime = json.loads((OUT/'baseline_runtime.json').read_text(encoding='utf-8'))
checks.update({p:digest(p)==sha for p,sha in runtime['code_sha256'].items()})
assert all(checks.values())
truth = json.loads((OUT/'ground_truth.json').read_text(encoding='utf-8'))
wb = openpyxl.load_workbook(OUT.parents[1]/'validation_003355_003718_manual.xlsx',data_only=True)
sheet = wb['검수 정답지']
assert sheet.max_row == 387
for row,record in enumerate(truth,2):
    assert str(sheet.cell(row,1).value).zfill(6)==record['image_id']
    assert sheet.cell(row,3).value==record['truth']
assert digest(OUT/'experiments.py')==digest(OLD/'experiments.py')
logs = [json.loads(s) for s in (OUT/'baseline.jsonl').read_text(encoding='utf-8').splitlines()]
assert len(logs)== 364 and {r['image_id'] for r in logs}==set(manifest['images'])
assert digest(OUT/'baseline.jsonl')==digest(OLD/'baseline.jsonl')
weights = {str(f):digest(f) for name in ('PP-OCRv5_mobile_det','PP-OCRv6_small_det','korean_PP-OCRv5_mobile_rec') for f in (ROOT/'weights/paddle'/name).glob('inference.*')}
assert len(weights)==9, 'Expected three files for each of three local models'
record = {'checked_at_utc':datetime.now(timezone.utc).isoformat(), 'verified_hash_count':len(checks),
          'all_source_image_code_hashes_match':all(checks.values()),'xlsx_truth_matches_snapshot':True,
          'baseline_reused_from':str(OLD),'baseline_log_sha256':digest(OUT/'baseline.jsonl'),
          'baseline_runtime_sha256':digest(OUT/'baseline_runtime.json'),
          'experiment_sha256':digest(OUT/'experiments.py'),'weights_current_sha256':weights,
          'limitations':['Baseline has no historical weight hashes. Current weight hashes are recorded, but historical weight identity cannot be proven.',
                          'Interrupted earlier p3 run is excluded from the final end-to-end score and timing.']}
prior = OUT/'provenance.json'
if prior.exists():
    initial=json.loads(prior.read_text(encoding='utf-8'))
    assert initial['weights_current_sha256']==weights
    record['weights_unchanged_since_first_check']=True
    prior=OUT/'provenance_final.json'
prior.write_text(json.dumps(record,ensure_ascii=False,indent=2),encoding='utf-8')
print({k:v for k,v in record.items() if k!='weights_current_sha256'})
