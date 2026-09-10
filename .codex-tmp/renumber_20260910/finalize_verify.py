from __future__ import annotations

import csv
import hashlib
import json
import math
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from statistics import median

from openpyxl import load_workbook
from PIL import Image, ImageOps

ROOT = Path(r"C:\ITDA_OCR_CODE")
IMAGES = ROOT / "추가수집데이터"
METADATA = IMAGES / "metadata"
LABELS = ROOT / "labels"
ARTIFACT = ROOT / "artifacts" / "renumber_additional_20260910"
FINAL_IDS = [f"{number:06d}" for number in range(3355, 3719)]
DELETED = [
    "003360", "003376", "003395", "003430", "003441", "003491",
    "003516", "003535", "003609", "003675", "003677", "003689",
    "003708", "003713", "003716", "003719", "003724", "003727",
    "003729", "003730", "003738", "003739",
]


def normalize_cell(value) -> str:
    if isinstance(value, (date, datetime)):
        return value.strftime("%Y-%m-%d")
    if value is None:
        return ""
    return str(value)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def write_json(path: Path, value) -> None:
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def update_paired_timing() -> None:
    for path in LABELS.rglob("paired_timing.json"):
        if "validation_000001_003352" in path.parts:
            continue
        data = json.loads(path.read_text(encoding="utf-8"))
        details = data.get("details", [])
        if not details:
            continue
        data["pairs"] = len(details)
        data["same_text_count"] = sum(bool(row.get("same_text")) for row in details)
        data["same_full_ocr_lines_count"] = sum(bool(row.get("same_lines")) for row in details)
        ratios = [row["ratio"] for row in details if row.get("ratio") is not None and math.isfinite(row["ratio"])]
        data["median_p3_over_baseline_ratio"] = median(ratios)
        data["baseline_original_seconds"] = sum(row.get("baseline_seconds", 0) for row in details)
        data["p3_original_seconds"] = sum(row.get("p3_seconds", 0) for row in details)
        write_json(path, data)


def update_manifests() -> None:
    with (METADATA / "source_manifest.csv").open(encoding="utf-8-sig", newline="") as handle:
        source_rows = list(csv.DictReader(handle))
    photo_data = {
        Path(row["new_filename"]).stem: {
            "path": str((IMAGES / row["new_filename"]).resolve()),
            "sha256": row["sha256"],
        }
        for row in source_rows
    }
    for path in LABELS.rglob("manifest.json"):
        if "validation_000001_003352" in path.parts:
            continue
        data = json.loads(path.read_text(encoding="utf-8"))
        if "images" not in data:
            continue
        data["images"] = {image_id: photo_data[image_id] for image_id in FINAL_IDS}
        source_names = [Path(item).name for item in data.get("sources", {})]
        sources = {}
        for name in source_names:
            candidates = list(LABELS.rglob(name))
            candidates = [candidate for candidate in candidates if "node_modules" not in candidate.parts]
            if not candidates:
                continue
            preferred = next((candidate for candidate in candidates if candidate.parent == LABELS), candidates[0])
            sources[str(preferred.resolve())] = sha256(preferred)
        data["sources"] = sources
        write_json(path, data)
        integrity = {source: sha256(Path(source)) == digest for source, digest in sources.items()}
        write_json(path.parent / "integrity.json", integrity)


def rebuild_report(target: Path, score_path: Path, comparison_path: Path) -> None:
    scores = json.loads(score_path.read_text(encoding="utf-8"))
    comparisons = json.loads(comparison_path.read_text(encoding="utf-8")) if comparison_path.exists() else {}
    lines = [
        "# 소비기한 OCR 364장 검증 기록",
        "",
        "2026-09-10에 중복 사진 22장을 삭제하고, 남은 사진과 라벨을 003355–003718로 재번호화했다. 아래 수치는 삭제된 레코드를 제외한 기존 OCR 로그를 다시 집계한 값이다. OCR 추론은 다시 실행하지 않았다.",
        "",
        "## 데이터 범위",
        "",
        "- 사진: 364장",
        "- 일련번호: 003355–003718",
        "- 삭제: 22장",
        "- 라벨 및 로그: 삭제된 행 제거 후 동일 순서로 재번호화",
        "",
        "## 단계별 재집계",
        "",
        "| 단계 | 전체 정답 | 완전 날짜 | 부분 날짜 | NONE | 이미지 시간 중앙값 | p95 |",
        "|---|---:|---:|---:|---:|---:|---:|",
    ]
    for stage, score in scores.items():
        categories = score["categories"]
        def cell(kind):
            item = categories[kind]
            return f'{item["correct"]}/{item["total"]}'
        lines.append(
            f'| {stage} | {score["correct"]}/{score["total"]} ({100*score["accuracy"]:.2f}%) | '
            f'{cell("full")} | {cell("partial")} | {cell("none")} | '
            f'{score["median_image_seconds"]:.3f}초 | {score["p95_image_seconds"]:.3f}초 |'
        )
    if comparisons:
        lines += ["", "## 단계 간 정답 변화", ""]
        for name, values in comparisons.items():
            lines.append(f'- {name}: 정답으로 변경 {len(values.get("gained", []))}건, 오답으로 변경 {len(values.get("lost", []))}건')
    lines += [
        "",
        "실행 환경과 전체 벽시계 시간 등 원래 386장 실행에만 해당하는 값은 재집계 표에서 제외했다. 행별 예측, 정답, 오류 유형, 처리시간과 OCR 증거는 보존했다.",
        "",
    ]
    target.write_text("\n".join(lines), encoding="utf-8")


def rebuild_reports() -> None:
    current = LABELS / "audit" / "validation_003355_003718_20260910_year_unrestricted"
    standard = LABELS / "audit" / "validation_003355_003718_20260910"
    rebuild_report(
        LABELS / "validation_003355_003718_report_20260910.md",
        current / "scores.json",
        current / "comparisons.json",
    )
    previous_report = current / "previous_results" / "validation_003355_003718_report_20260910.md"
    if previous_report.exists():
        rebuild_report(previous_report, standard / "scores.json", standard / "comparisons.json")


def normalize_workbooks() -> None:
    workbooks = [
        path for path in LABELS.rglob("validation_003355_003718*.xlsx")
        if "node_modules" not in path.parts
    ]
    assert len(workbooks) == 5
    for workbook_path in workbooks:
        workbook = load_workbook(workbook_path, read_only=False, data_only=False)
        sheet = workbook["검수 정답지"]
        if sheet.max_row > 365:
            sheet.delete_rows(366, sheet.max_row - 365)
        workbook.save(workbook_path)
        values = []
        formulas = []
        for row_number in range(1, 366):
            row_values = []
            row_formulas = []
            for column in range(1, 9):
                value = sheet.cell(row_number, column).value
                if isinstance(value, str) and value.startswith("="):
                    row_formulas.append(value)
                    if column == 4 and row_number >= 2:
                        left = sheet.cell(row_number, 2).value
                        right = sheet.cell(row_number, 3).value
                        row_values.append("" if left in (None, "") or right in (None, "") else "True" if left == right else "False")
                    else:
                        row_values.append("")
                else:
                    row_values.append("" if value is None else value)
                    row_formulas.append("")
            values.append(row_values)
            formulas.append(row_formulas)
        lines = [
            json.dumps({"kind": "workbook", "source": str(workbook_path.resolve()),
                        "sha256": sha256(workbook_path), "sheets": 1, "rows": 365,
                        "dataRows": 364, "columns": 8}, ensure_ascii=False),
            json.dumps({"kind": "sheet", "name": sheet.title, "range": "A1:H365"}, ensure_ascii=False),
        ]
        for index, (row_values, row_formulas) in enumerate(zip(values, formulas), 1):
            lines.append(json.dumps({"kind": "row", "sheet": sheet.title, "row": index,
                                     "values": row_values, "formulas": row_formulas}, ensure_ascii=False))
        Path(str(workbook_path) + ".inspect.ndjson").write_text("\n".join(lines) + "\n", encoding="utf-8")


def csv_rows(path: Path):
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.reader(handle))


def verify_workbooks() -> list[str]:
    checked = []
    workbooks = [
        path for path in LABELS.rglob("validation_003355_003718*.xlsx")
        if "node_modules" not in path.parts
    ]
    assert len(workbooks) == 5
    for workbook_path in workbooks:
        workbook = load_workbook(workbook_path, read_only=False, data_only=False)
        assert workbook.sheetnames == ["검수 정답지"]
        sheet = workbook["검수 정답지"]
        assert sheet.max_row >= 365 and sheet.max_column == 8
        assert len(sheet.tables) == 1
        table = next(iter(sheet.tables.values()))
        assert table.ref == "A1:H365"
        values = list(sheet.iter_rows(min_row=1, max_row=365, max_col=8, values_only=True))
        trailing = list(sheet.iter_rows(min_row=366, max_row=sheet.max_row, max_col=8, values_only=True))
        assert all(all(value is None for value in row) for row in trailing)
        assert [int(row[0]) for row in values[1:]] == list(range(3355, 3719))
        for index, xlsx_row in enumerate(values[1:], 2):
            assert sheet.cell(index, 4).value == f'=IF(OR(B{index}="",C{index}=""),"",IF(B{index}=C{index},"True","False"))'
        inspect_path = Path(str(workbook_path) + ".inspect.ndjson")
        inspect_lines = [json.loads(line) for line in inspect_path.read_text(encoding="utf-8").splitlines() if line]
        assert inspect_lines[0]["dataRows"] == 364
        assert inspect_lines[0]["sha256"] == sha256(workbook_path)
        assert len([row for row in inspect_lines if row.get("kind") == "row"]) == 365
        checked.append(str(workbook_path.relative_to(ROOT)))
    return checked


def verify_all() -> dict:
    photos = sorted(IMAGES.glob("*.jpg"))
    assert [path.stem for path in photos] == FINAL_IDS
    assert not list(IMAGES.glob(".__renumber__*"))
    with (METADATA / "source_manifest.csv").open(encoding="utf-8-sig", newline="") as handle:
        manifest = list(csv.DictReader(handle))
    assert [Path(row["new_filename"]).stem for row in manifest] == FINAL_IDS
    assert all(sha256(IMAGES / row["new_filename"]) == row["sha256"] for row in manifest)
    annotations = json.loads((METADATA / "annotations.json").read_text(encoding="utf-8"))
    assert list(annotations) == [f"{image_id}.jpg" for image_id in FINAL_IDS]
    for photo in (photos[0], photos[len(photos)//2], photos[-1]):
        with Image.open(photo) as source:
            image = ImageOps.exif_transpose(source)
            record = annotations[photo.name]
            assert (record["width"], record["height"]) == image.size

    csv_files = [
        path for path in LABELS.rglob("*.csv")
        if "node_modules" not in path.parts
        and "validation_000001_003352" not in path.parts
        and not path.name.startswith("validation_000001_003352")
    ]
    checked_csv = 0
    for path in csv_files:
        rows = csv_rows(path)
        if rows and rows[0] and rows[0][0] == "image_id":
            assert len(rows) == 365, (path, len(rows))
            assert [int(row[0]) for row in rows[1:]] == list(range(3355, 3719)), path
            checked_csv += 1

    full_logs = 0
    partial_logs = 0
    for path in LABELS.rglob("*.jsonl"):
        if "validation_000001_003352" in path.parts:
            continue
        rows = [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line]
        ids = [row["image_id"] for row in rows if isinstance(row, dict) and "image_id" in row]
        assert len(ids) == len(set(ids))
        assert set(ids) <= set(FINAL_IDS)
        if path.name == "p3.jsonl" and "20260909" in str(path):
            partial_logs += 1
        else:
            assert ids == FINAL_IDS, (path, len(ids), ids[:2], ids[-2:])
            full_logs += 1

    for path in LABELS.rglob("ground_truth.json"):
        if "validation_000001_003352" in path.parts:
            continue
        rows = json.loads(path.read_text(encoding="utf-8"))
        assert [row["image_id"] for row in rows] == FINAL_IDS
    for path in LABELS.rglob("scores.json"):
        if "validation_000001_003352" in path.parts:
            continue
        scores = json.loads(path.read_text(encoding="utf-8"))
        for score in scores.values():
            assert score["total"] == 364
            assert score["correct"] + score["incorrect"] == 364
            assert sum(group["total"] for group in score["categories"].values()) == 364

    assert not list(LABELS.rglob("*003355_003740*"))
    workbooks = verify_workbooks()
    return {
        "deleted_photos": 22,
        "remaining_photos": len(photos),
        "number_range": "003355-003718",
        "metadata_rows": len(manifest),
        "annotation_records": len(annotations),
        "label_csv_files_checked": checked_csv,
        "full_jsonl_logs_checked": full_logs,
        "partial_jsonl_logs_checked": partial_logs,
        "xlsx_files_checked": len(workbooks),
        "xlsx_paths": workbooks,
    }


def main() -> None:
    update_paired_timing()
    update_manifests()
    rebuild_reports()
    normalize_workbooks()
    update_manifests()
    result = verify_all()
    write_json(ARTIFACT / "검증결과.json", result)
    report = [
        "# 추가수집데이터 중복 삭제 및 재번호화 완료",
        "",
        "- 삭제한 사진: 22장",
        "- 남은 사진: 364장",
        "- 최종 번호: 003355–003718",
        f'- 갱신·검증한 라벨 CSV: {result["label_csv_files_checked"]}개',
        f'- 갱신·검증한 전체 JSONL 로그: {result["full_jsonl_logs_checked"]}개',
        f'- 갱신·검증한 XLSX: {result["xlsx_files_checked"]}개',
        "",
        "사진 파일명, annotations.json, source_manifest.csv, 라벨 CSV·JSON·JSONL·Markdown·XLSX 및 inspect 기록에 같은 번호 매핑을 적용했다.",
        "",
        "번호별 변경 내역은 `번호변경표.csv`, 삭제 목록은 `삭제한_사진_22장.txt`에 기록했다.",
        "",
    ]
    (ARTIFACT / "완료보고서.md").write_text("\n".join(report), encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
