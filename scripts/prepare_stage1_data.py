"""Freeze approved date labels and prepare (not approve) the OCR annotation queue.

Run prepare, edit the four XLSX files with the staged artifact-tool builder,
then run finalize. verify rechecks the frozen outputs without modifying them.
"""
import argparse
import csv
import hashlib
import json
import re
import shutil
from collections import Counter, defaultdict
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timezone
from pathlib import Path

import openpyxl  # Read/verify only. XLSX authoring uses artifact-tool.
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / '학습 및 테스트 결과'
TMP = OUT / 'tmp/stage1'
DATA = OUT / '00_protocol/data'
BACKUP = DATA / 'originals'
BOOKS = {
    'product': ('상품사진_정답지/answer_000001_003352_manual.xlsx', 3352),
    'additional': ('추가수집_정답지/answer_003353_003716_manual.xlsx', 364),
    'test': ('테스트용_정답지/answer_000001_003716_manual.xlsx', 3716),
    'training': ('학습대상_정답지/answer_000001_002610_manual.xlsx', 2610),
}
MAPPING = 'artifacts/training_dataset_20260910/source_mapping.json'
EXCLUDED = 'artifacts/training_dataset_20260910/excluded_640x640.json'
ADDMAN = '추가수집데이터/metadata/source_manifest.csv'
TESTMAN = '테스트용데이터/metadata/source_manifest.csv'
ANNOTATIONS = '추가수집데이터/metadata/annotations.json'
LEGACY = 'labels/validation_000001_003352_manual.csv'
DEVELOPMENT = 'artifacts/date-recognition-repair-20260910/manifest.json'


def digest(path):
    with path.open('rb') as stream:
        h = hashlib.sha256()
        for chunk in iter(lambda: stream.read(1024 * 1024), b''):
            h.update(chunk)
    return h.hexdigest()


def read_json(path):
    return json.loads(path.read_text(encoding='utf-8-sig'))


def write_json(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')


def write_jsonl(path, rows):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(''.join(json.dumps(r, ensure_ascii=False) + '\n' for r in rows), encoding='utf-8')


def read_csv(path):
    with path.open(encoding='utf-8-sig', newline='') as stream:
        return list(csv.DictReader(stream))


def write_csv(path, rows, fields=None):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open('w', encoding='utf-8-sig', newline='') as stream:
        writer = csv.DictWriter(stream, fieldnames=fields or list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def backup(relative):
    target = BACKUP / relative
    target.parent.mkdir(parents=True, exist_ok=True)
    if not target.exists():
        shutil.copy2(ROOT / relative, target)
    return target


def book_rows(path):
    wb = openpyxl.load_workbook(path, data_only=False)
    assert wb.sheetnames == ['검수 정답지'], path
    rows = []
    for row in wb.active.iter_rows(min_row=2, max_col=8):
        if row[0].value is None:
            continue
        key = str(row[0].value).zfill(6)
        assert re.fullmatch(r'\d{6}', key), (path, key)
        rows.append(dict(id=key, row=row[0].row, values=[c.value for c in row]))
    assert len({r['id'] for r in rows}) == len(rows), path
    wb.close()
    return rows


def date_fields(value):
    from src.date_extraction import submission_fields
    assert value != 'NONE-NONE-NONE', value
    fields = submission_fields(value)
    assert fields['final_date'] == value, value
    return fields


def inspect_image(path):
    with Image.open(path) as im:
        width, height = im.size
        im.verify()
    return dict(path=path.relative_to(ROOT).as_posix(), sha256=digest(path),
                size_bytes=path.stat().st_size, width=width, height=height)


def current_images():
    paths = sorted(p for folder in ('상품사진입니다', '추가수집데이터', '테스트용데이터', '학습대상데이터')
                   for p in (ROOT / folder).iterdir() if p.suffix.lower() in ('.jpg', '.jpeg', '.png'))
    with ThreadPoolExecutor(max_workers=4) as pool:
        return list(pool.map(inspect_image, paths))


def prepare():
    DATA.mkdir(parents=True, exist_ok=True)
    if (DATA / 'frozen_hashes.json').exists():
        for relative, expected in read_json(DATA / 'frozen_hashes.json').items():
            assert digest(ROOT / relative) == expected, ('changed_since_freeze', relative)
    originals = {p: digest(backup(p)) for p in [*(p for p, _ in BOOKS.values()),
                 MAPPING, EXCLUDED, ADDMAN, TESTMAN, ANNOTATIONS, LEGACY, DEVELOPMENT]}
    write_json(DATA / 'original_hashes.json', originals)
    images = current_images()
    sources = [r for r in images if r['path'].split('/')[0] in ('상품사진입니다', '추가수집데이터')]
    assert len(sources) == 3716
    hashes = defaultdict(list)
    for item in sources:
        hashes[item['sha256']].append(item)
    mapping = read_json(BACKUP / MAPPING)
    repaired = []
    repairs = []
    for row in mapping:
        matches = hashes[row['sha256']]
        assert len(matches) == 1, row
        updated = dict(row, source=matches[0]['path'])
        if row['source'].replace('\\', '/') != updated['source']:
            repairs.append(dict(file=MAPPING, key=row['filename'], old=row['source'], new=updated['source'], sha256=row['sha256']))
        repaired.append(updated)
    assert len(repaired) == 2610
    training_sources = {Path(r['filename']).stem[4:]: Path(r['source']).stem for r in repaired}
    legacy = {r['image_id'].zfill(6): r for r in read_csv(BACKUP / LEGACY)}
    all_books = {key: book_rows(BACKUP / p) for key, (p, _) in BOOKS.items()}
    dates = {}
    for key in ('product', 'additional'):
        for row in all_books[key]:
            value = str(row['values'][2]).strip()
            value = '2027-07-08' if row['id'] == '000157' else 'NONE' if value == 'NONE-NONE-NONE' else value
            date_fields(value)
            dates[row['id']] = value
    additional_manifest = read_csv(BACKUP / ADDMAN)
    old_to_new = {}
    for row in additional_manifest:
        match = hashes[row['sha256']]
        assert len(match) == 1
        new = Path(match[0]['path']).name
        old_to_new[row['new_filename']] = new
        if row['new_filename'] != new:
            repairs.append(dict(file=ADDMAN, key=row['new_filename'], old=row['new_filename'], new=new, sha256=row['sha256']))
        row['new_filename'] = new
    annotation_source = read_json(BACKUP / ANNOTATIONS)
    assert set(annotation_source) == set(old_to_new)
    annotations = {old_to_new[k]: v for k, v in annotation_source.items()}
    assert len(annotations) == 364
    for old, new in old_to_new.items():
        if old != new:
            repairs.append(dict(file=ANNOTATIONS, key=old, old=old, new=new))
    test_manifest = read_csv(BACKUP / TESTMAN)
    for row in test_manifest:
        match = hashes[row['sha256']]
        assert len(match) == 1
        new = match[0]['path']
        if row['source_path'].replace('\\', '/') != new:
            repairs.append(dict(file=TESTMAN, key=row['serial_number'], old=row['source_path'], new=new, sha256=row['sha256']))
        row['source_path'] = new
    metadata = {}
    for key, value in dates.items():
        old = legacy.get(key, {})
        tags = [t.strip() for t in old.get('오류 유형', '').split(';') if t.strip()]
        # Existing tags are imported evidence, not a new human image assessment.
        kind = 'no_expiry' if value == 'NONE' else 'partial_date' if 'NONE' in value else 'full_date'
        tags = list(dict.fromkeys([*tags, kind]))
        annotation = annotations.get(key + '.jpg')
        if annotation:
            lines = [a for a in annotation['ann'] if a['cls'] in ('date', 'exp')]
            if len(lines) > 1:
                tags.append('multiple_date_lines')
            if any(re.search('[A-Za-z]', a.get('transcription', '')) for a in lines):
                tags.append('latin_text')
        metadata[key] = dict(difficulty=old.get('난이도') or 'unassessed',
                             difficulty_status='legacy_imported' if old.get('난이도') else 'needs_review',
                             condition_tags=tags, condition_tags_status='pending_human_review',
                             legacy_label_status=old.get('라벨 상태') or None)
    plan = []
    for key, (relative, count) in BOOKS.items():
        rows = all_books[key]
        assert len(rows) == count
        changes, meta = [], []
        corrected_row = None
        for row in rows:
            sid = training_sources[row['id']] if key == 'training' else row['id']
            value = dates[sid]
            old_value = row['values'][2]
            assert str(old_value).strip() == value or (old_value == 'NONE-NONE-NONE' and value == 'NONE') or sid == '000157'
            if old_value != value:
                changes.append(dict(cell=f"C{row['row']}", old=old_value, value=value))
            if sid == '000157':
                corrected_row = row['row']
            item = metadata[sid]
            note = '최종 날짜만 승인; OCR 주석 별도 검수'
            if sid == '000157':
                note += '; 사용자 확정 2027-07-08'
            meta.append(['approved', item['difficulty'], ';'.join(item['condition_tags']), note])
        assert [r['row'] for r in rows] == list(range(2, count + 2))
        plan.append(dict(key=key, path=relative, count=count, dateChanges=changes,
                         correctedRow=corrected_row, metadata=meta))
    write_json(TMP / 'workbook_edits.json', plan)
    write_json(TMP / 'prepared_data.json', dict(images=images, sources=sources, mapping=repaired,
               additional_manifest=additional_manifest, test_manifest=test_manifest,
               annotations=annotations, repairs=repairs, dates=dates, metadata=metadata))
    print(json.dumps(dict(images=len(images), sources=len(sources), training=len(repaired),
                         path_repairs=dict(Counter(r['file'] for r in repairs)),
                         workbook_date_changes={e['key']:len(e['dateChanges']) for e in plan}), ensure_ascii=False))


def install_workbooks(plan):
    # Check all staged files before replacing any of the four user workbooks.
    for entry in plan:
        old = openpyxl.load_workbook(BACKUP / entry['path'], data_only=False)
        new = openpyxl.load_workbook(TMP / (entry['key'] + '.xlsx'), data_only=False)
        assert old.sheetnames == new.sheetnames
        expected_changes = {c['cell']: c['value'] for c in entry['dateChanges']}
        for index, values in enumerate(entry['metadata'], 2):
            expected_changes.update({f'{col}{index}': v for col, v in zip('EFGH', values)})
        for row in old.active.iter_rows():
            for cell in row:
                expected = expected_changes.get(cell.coordinate, cell.value)
                actual = new.active[cell.coordinate]
                assert actual.value == expected, (entry['key'], cell.coordinate, expected, actual.value)
                assert cell.number_format == actual.number_format, (entry['key'], cell.coordinate, 'number_format')
        assert old.active.freeze_panes == new.active.freeze_panes
        assert str(old.active.merged_cells) == str(new.active.merged_cells)
        assert list(old.active.tables) == list(new.active.tables)
        assert old.active.auto_filter.ref == new.active.auto_filter.ref
        assert len(old.active.conditional_formatting) == len(new.active.conditional_formatting)
        assert len(old.active.data_validations.dataValidation) == len(new.active.data_validations.dataValidation)
        old.close()
        new.close()
    for entry in plan:
        shutil.copy2(TMP / (entry['key'] + '.xlsx'), ROOT / entry['path'])


def finalize():
    plan = read_json(TMP / 'workbook_edits.json')
    prepared = read_json(TMP / 'prepared_data.json')
    install_workbooks(plan)
    write_json(ROOT / MAPPING, prepared['mapping'])
    write_csv(ROOT / ADDMAN, prepared['additional_manifest'])
    write_csv(ROOT / TESTMAN, prepared['test_manifest'])
    write_json(ROOT / ANNOTATIONS, prepared['annotations'])
    write_json(DATA / 'provenance_repairs.json', prepared['repairs'])
    books = {key: dict(path=p, sha256=digest(ROOT / p), rows=n) for key, (p, n) in BOOKS.items()}
    images = {r['path']: r for r in prepared['images']}
    training = {r['source']: r for r in prepared['mapping']}
    excluded = {p.replace('\\', '/') for p in read_json(ROOT / EXCLUDED)}
    development_hashes = {r['sha256'] for r in read_json(ROOT / DEVELOPMENT)['included']}
    inventory, queue, lines = [], [], []
    now = datetime.now(timezone.utc).isoformat()
    for source in prepared['sources']:
        sid = Path(source['path']).stem
        product = source['path'].startswith('상품사진입니다/')
        source_book = books['product' if product else 'additional']
        source_row = int(sid) + 1 if product else int(sid) - 3353 + 2
        test_stem = ('AMLT' if product else 'BMLT') + sid
        test_path = '테스트용데이터/' + test_stem + Path(source['path']).suffix
        copies = [dict(kind='source', **source), dict(kind='test', **images[test_path])]
        train = training.get(source['path'])
        if train:
            copies.append(dict(kind='training', **images['학습대상데이터/' + train['filename']]))
        assert all(r['sha256'] == source['sha256'] for r in copies)
        meta = prepared['metadata'][sid]
        final = prepared['dates'][sid]
        annotation = prepared['annotations'].get(sid + '.jpg') if not product else None
        record = dict(asset_id='sha256:' + source['sha256'], source_dataset='product' if product else 'additional',
            source_image_id=sid, source_path=source['path'], test_image_id=test_stem,
            training_image_id=Path(train['filename']).stem if train else None,
            selected_for_training=bool(train), final_date=final, **{k:v for k,v in date_fields(final).items() if k != 'final_date'},
            final_date_status='approved', approved_by='user', approval_recorded_at=now,
            approval_basis='User-designated XLSX final dates; 000157 explicitly confirmed as 2027-07-08; missing sentinel NONE',
            label_source=dict(**source_book, sheet='검수 정답지', cell=f'C{source_row}'),
            copies=copies, **meta, group_id='sha256:' + source['sha256'],
            group_status='exact_identity_only_pending_product_review',
            seen_in_development=source['sha256'] in development_hashes,
            fold5_eligible_after_group_review=bool(train) and source['sha256'] not in development_hashes,
            exclusion_reason='preexisting_640x640_augmented_derivative' if source['path'] in excluded else None,
            recognition_transcription_status='needs_review' if annotation else 'not_annotated',
            detection_polygon_status='needs_review' if annotation else 'not_annotated',
            approved_for_recognition_training=False, approved_for_detection_training=False)
        inventory.append(record)
        if not train:
            continue
        annotations = []
        if annotation:
            assert (annotation['width'], annotation['height']) == (source['width'], source['height'])
            for index, item in enumerate(annotation['ann']):
                if item['cls'] not in ('date', 'exp'):
                    continue
                x1,y1,x2,y2 = item['bbox']
                issues = []
                if not (0 <= x1 < x2 <= source['width'] and 0 <= y1 < y2 <= source['height']):
                    issues.append('bbox_out_of_bounds')
                if not item.get('transcription', '').strip():
                    issues.append('empty_transcription')
                line = dict(annotation_id=f'{record["training_image_id"]}:source:{index}',
                    asset_id=record['asset_id'], image_id=record['training_image_id'], source_image_id=sid,
                    source_path=source['path'], training_path='학습대상데이터/' + train['filename'],
                    transcription=item.get('transcription'), bbox=item['bbox'],
                    polygon=[[x1,y1],[x2,y1],[x2,y2],[x1,y2]],
                    polygon_origin='source_bbox_rectangle', source_class=item['cls'],
                    role='unknown', selected_for_final_date=None,
                    status='needs_review', automatic_check='passed' if not issues else 'failed',
                    automatic_issues=issues, approved_by=None, approved_at=None,
                    source_annotation_path=ANNOTATIONS, source_annotation_key=sid+'.jpg', source_annotation_index=index)
                annotations.append(line['annotation_id'])
                lines.append(line)
        queue.append(dict(image_id=record['training_image_id'], asset_id=record['asset_id'],
            source_image_id=sid, image_path='학습대상데이터/' + train['filename'], final_date=final,
            final_date_status='approved', annotation_ids=annotations,
            action='review_source_annotations_and_check_missing_lines' if annotation else 'create_date_line_annotations',
            recognition_status=record['recognition_transcription_status'], detection_status=record['detection_polygon_status'],
            group_id=record['group_id'], group_status=record['group_status'],
            difficulty=record['difficulty'], difficulty_status=record['difficulty_status'],
            condition_tags=record['condition_tags'], condition_tags_status=record['condition_tags_status'],
            seen_in_development=record['seen_in_development'], review_status='pending',
            note='NONE final date does not prove absence of manufacturing/date regions; inspect all date lines',
            approved_by=None, approved_at=None))
    assert len(excluded) == 1106
    assert excluded == {r['path'] for r in prepared['sources'] if (r['width'],r['height']) == (640,640)}
    write_json(OUT / '00_protocol/dataset_inventory.json', dict(schema_version=1, created_at=now,
        missing_final_date='NONE', workbooks=books, source_assets=len(inventory), training_assets=len(queue), images=inventory))
    write_jsonl(DATA / 'source_labels.jsonl', inventory)
    train_records = [r for r in inventory if r['selected_for_training']]
    write_jsonl(DATA / 'training_labels.jsonl', train_records)
    for key in BOOKS:
        rows = book_rows(ROOT / BOOKS[key][0])
        mapping_ids = {Path(r['filename']).stem[4:]: Path(r['filename']).stem for r in prepared['mapping']}
        output = []
        for row in rows:
            values = row['values']
            output_id = mapping_ids[row['id']] if key == 'training' else (('AMLT' if int(row['id']) <= 3352 else 'BMLT') + row['id']) if key == 'test' else row['id']
            output.append(dict(zip(['image_id','추출한 날짜','정답 날짜','True/False','라벨 상태','난이도','오류 유형','비고'],
                                  [output_id, values[1] or '', values[2], '', *values[4:]])))
        write_csv(DATA / f'{key}_labels.csv', output)
    write_jsonl(DATA / 'annotation_review_queue.jsonl', queue)
    write_jsonl(DATA / 'source_annotation_candidates.jsonl', lines)
    write_jsonl(DATA / 'excluded_images.jsonl', [r for r in inventory if r['exclusion_reason']])
    write_json(DATA / 'annotation_checks.json', dict(images=364, date_lines=len(lines),
        automatic_passed=sum(r['automatic_check']=='passed' for r in lines),
        issues=[r for r in lines if r['automatic_issues']], human_approved=0,
        check_scope='SHA linkage, image dimensions, bbox order/bounds, nonempty transcription; raw text is not forced to ISO',
        source_sha256=digest(ROOT / ANNOTATIONS)))
    write_json(DATA / 'workbook_changes.json', [{k:v for k,v in r.items() if k!='metadata'} for r in plan])
    frozen = [p for p in DATA.iterdir() if p.is_file() and p.name != 'frozen_hashes.json']
    frozen += [OUT / '00_protocol/dataset_inventory.json', *(ROOT / p for p,_ in BOOKS.values()),
               *(ROOT / p for p in (MAPPING, ADDMAN, TESTMAN, ANNOTATIONS, EXCLUDED))]
    write_json(DATA / 'frozen_hashes.json', {p.relative_to(ROOT).as_posix(): digest(p) for p in frozen})
    verify(write_report=True)


def verify(write_report=False):
    inventory = read_json(OUT / '00_protocol/dataset_inventory.json')
    records = inventory['images']
    assert len(records) == 3716
    assert len({r['source_image_id'] for r in records}) == 3716
    selected = [r for r in records if r['selected_for_training']]
    assert len(selected) == len({r['training_image_id'] for r in selected}) == 2610
    for relative, expected in read_json(DATA / 'frozen_hashes.json').items():
        assert digest(ROOT / relative) == expected, relative
    for relative, expected in read_json(DATA / 'original_hashes.json').items():
        assert digest(BACKUP / relative) == expected, ('backup', relative)
    current = {r['path']:r for r in current_images()}
    assert len(current) == 10042
    for row in records:
        date_fields(row['final_date'])
        assert row['final_date_status'] == 'approved'
        for copy in row['copies']:
            assert current[copy['path']] == {k:v for k,v in copy.items() if k!='kind'}, copy['path']
        assert row['difficulty'] and row['condition_tags']
    for entry in read_json(TMP / 'workbook_edits.json'):
        rows = book_rows(ROOT / entry['path'])
        assert len(rows) == entry['count']
        source_by_id = {r['source_image_id']:r for r in records}
        training_by_id = {r['training_image_id'][4:]:r for r in selected}
        for row in rows:
            expected = training_by_id[row['id']] if entry['key']=='training' else source_by_id[row['id']]
            assert row['values'][2] == expected['final_date'], (entry['key'], row['id'])
            assert row['values'][4:8] == entry['metadata'][row['row']-2], (entry['key'], row['id'], 'metadata')
        for change in entry['dateChanges']:
            row = next(r for r in rows if r['row']==int(change['cell'][1:]))
            assert row['values'][2] == change['value']
    for row in read_json(ROOT / MAPPING):
        assert digest(ROOT / row['source']) == row['sha256'] == digest(ROOT / '학습대상데이터' / row['filename'])
    for row in read_csv(ROOT / ADDMAN):
        assert digest(ROOT / '추가수집데이터' / row['new_filename']) == row['sha256']
    for row in read_csv(ROOT / TESTMAN):
        assert digest(ROOT / row['source_path']) == row['sha256'] == digest(ROOT / '테스트용데이터' / row['dataset_path'])
    queue = [json.loads(r) for r in (DATA / 'annotation_review_queue.jsonl').read_text(encoding='utf-8').splitlines()]
    candidates = [json.loads(r) for r in (DATA / 'source_annotation_candidates.jsonl').read_text(encoding='utf-8').splitlines()]
    assert len(queue) == len({r['image_id'] for r in queue}) == 2610
    assert len(candidates) == len({r['annotation_id'] for r in candidates}) == 404
    annotations = read_json(ROOT / ANNOTATIONS)
    for candidate in candidates:
        original = annotations[candidate['source_annotation_key']]['ann'][candidate['source_annotation_index']]
        assert candidate['bbox'] == original['bbox'] and candidate['transcription'] == original['transcription']
        assert candidate['status'] == 'needs_review' and candidate['approved_by'] is None
    assert all(not r['approved_for_recognition_training'] and not r['approved_for_detection_training'] for r in records)
    result = dict(status='passed', stage_1_complete=True, ready_for_stage_2=True,
        verified_at=datetime.now(timezone.utc).isoformat(), source_images=3716, physical_image_files=10042,
        training_images=2610, approved_final_dates=2610,
        training_date_kinds=dict(Counter('none' if r['final_date']=='NONE' else 'partial' if 'NONE' in r['final_date'] else 'full' for r in selected)),
        corrected_source_000157=next(r['final_date'] for r in records if r['source_image_id']=='000157'),
        excluded_640x640=1106, existing_development_images=sum(r['seen_in_development'] for r in selected),
        duplicate_ids=0, missing_files=0, invalid_dates=0, provenance_hash_mismatches=0,
        stale_provenance_paths=0, missing_metadata_values=0,
        recognition_training_approved=0, detection_training_approved=0,
        source_annotation_candidates=read_json(DATA / 'annotation_checks.json'),
        stage_2_work=['Human review/create raw date-line transcription and polygons',
                      'Confirm product/sequence/derivative groups and unassessed difficulty/condition tags',
                      'Create group-safe folds and permanent inner validation only after review; 705 seen assets barred from fold 5'],
        training_allowed=False)
    if write_report:
        write_json(OUT / '00_protocol/stage1_readiness.json', result)
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    import sys
    sys.path.insert(0, str(ROOT))
    parser = argparse.ArgumentParser()
    parser.add_argument('mode', choices=('prepare', 'finalize', 'verify'))
    args = parser.parse_args()
    {'prepare': prepare, 'finalize': finalize, 'verify': verify}[args.mode]()
